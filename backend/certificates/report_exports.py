"""
Shared PDF/Excel export helpers for the Reports page. Every report view
builds a plain (headers, rows) table and hands it to one of these two
functions rather than each view reinventing file generation.

reportlab is used for PDF (not weasyprint) specifically because it's
pure Python with no system-level dependencies (no Cairo/Pango/GTK), so
it installs cleanly on Windows via pip with nothing extra to configure
— important since this app is developed and run locally on Windows.
"""
import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND_GREEN = '0B4D3B'


def build_xlsx_response(filename, title, headers, rows):
    """rows is a list of lists/tuples, same length/order as headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else 'Report'  # Excel sheet name limit

    ws.append(headers)
    header_fill = PatternFill(start_color=BRAND_GREEN, end_color=BRAND_GREEN, fill_type='solid')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill

    for row in rows:
        ws.append(list(row))

    # Rough auto-width based on the longest value seen in each column.
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_len = max(max_len, len(str(row[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def build_pdf_response(filename, title, headers, rows, subtitle=None):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        topMargin=36, bottomMargin=36, leftMargin=36, rightMargin=36,
    )
    styles = getSampleStyleSheet()

    elements = [Paragraph(title, styles['Title'])]
    if subtitle:
        elements.append(Paragraph(subtitle, styles['Normal']))
    elements.append(Spacer(1, 16))

    table_data = [headers] + [list(row) for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{BRAND_GREEN}')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(table)

    if not rows:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph('No data for this report.', styles['Normal']))

    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
